"""Export service — CSV and Excel report generation."""

import io
import csv
from datetime import datetime

from app.models import Complaint


def export_complaints_csv(depot_id, filters=None):
    """Export depot complaints as CSV. Returns a StringIO buffer."""
    complaints = _get_filtered_complaints(depot_id, filters)

    output = io.StringIO()
    writer = csv.writer(output)

    # Header
    writer.writerow([
        "Reference Number",
        "Category",
        "Bus",
        "Route",
        "Reported Date",
        "Reported Time",
        "Status",
        "Priority",
        "Created At",
        "Resolved At",
    ])

    for c in complaints:
        writer.writerow([
            c.reference_number,
            c.category,
            c.bus.bus_number if c.bus else "",
            f"{c.route.source} → {c.route.destination}" if c.route else "",
            c.reported_date.isoformat() if c.reported_date else "",
            c.reported_time.strftime("%H:%M") if c.reported_time else "",
            c.status,
            c.priority,
            c.created_at.strftime("%Y-%m-%d %H:%M") if c.created_at else "",
            c.resolved_at.strftime("%Y-%m-%d %H:%M") if c.resolved_at else "",
        ])

    output.seek(0)
    return output


def export_complaints_xlsx(depot_id, filters=None):
    """Export depot complaints as Excel. Returns a BytesIO buffer."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    complaints = _get_filtered_complaints(depot_id, filters)

    wb = Workbook()
    ws = wb.active
    ws.title = "Complaints Report"

    # Header styling
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1A237E", end_color="1A237E", fill_type="solid")

    headers = [
        "Reference Number",
        "Category",
        "Bus",
        "Route",
        "Reported Date",
        "Reported Time",
        "Status",
        "Priority",
        "Created At",
        "Resolved At",
    ]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row, c in enumerate(complaints, 2):
        ws.cell(row=row, column=1, value=c.reference_number)
        ws.cell(row=row, column=2, value=c.category)
        ws.cell(row=row, column=3, value=c.bus.bus_number if c.bus else "")
        ws.cell(row=row, column=4, value=f"{c.route.source} → {c.route.destination}" if c.route else "")
        ws.cell(row=row, column=5, value=c.reported_date.isoformat() if c.reported_date else "")
        ws.cell(row=row, column=6, value=c.reported_time.strftime("%H:%M") if c.reported_time else "")
        ws.cell(row=row, column=7, value=c.status)
        ws.cell(row=row, column=8, value=c.priority)
        ws.cell(row=row, column=9, value=c.created_at.strftime("%Y-%m-%d %H:%M") if c.created_at else "")
        ws.cell(row=row, column=10, value=c.resolved_at.strftime("%Y-%m-%d %H:%M") if c.resolved_at else "")

    # Auto-adjust column widths
    for col in ws.columns:
        max_length = 0
        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except Exception:
                pass
        adjusted_width = min(max_length + 2, 40)
        ws.column_dimensions[col[0].column_letter].width = adjusted_width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def _get_filtered_complaints(depot_id, filters=None):
    """Get complaints for export with optional filters."""
    query = Complaint.query.filter_by(depot_id=depot_id)

    if filters:
        if filters.get("status"):
            query = query.filter_by(status=filters["status"])
        if filters.get("category"):
            query = query.filter_by(category=filters["category"])
        if filters.get("date_from"):
            try:
                date_from = datetime.fromisoformat(filters["date_from"]).date()
                query = query.filter(Complaint.reported_date >= date_from)
            except ValueError:
                pass
        if filters.get("date_to"):
            try:
                date_to = datetime.fromisoformat(filters["date_to"]).date()
                query = query.filter(Complaint.reported_date <= date_to)
            except ValueError:
                pass

    return query.order_by(Complaint.created_at.desc()).all()
